"""
Contract Parser Module
Extracts and parses contract information from PDF, DOCX, and text files using Claude AI.
Supports JSON-only output format.
"""

import os
import json
from anthropic import Anthropic
from typing import Dict, Optional
from datetime import datetime
from pathlib import Path


class ContractParser:
    """Parse contract documents using Claude AI with JSON output."""

    def __init__(self, api_key: Optional[str] = None):
        """
        Initialize the contract parser.

        Args:
            api_key: Anthropic API key. If None, reads from ANTHROPIC_API_KEY env var.
        """
        self.api_key = api_key or os.getenv('ANTHROPIC_API_KEY')
        if not self.api_key:
            raise ValueError("ANTHROPIC_API_KEY must be set in environment or passed to constructor")

        self.client = Anthropic(api_key=self.api_key)
        self.model = "claude-sonnet-4-5"  # Updated to latest model

    def extract_text(self, sourc: str) -> str:
        """
        Extract text from PDF, DOCX, or plain text file.

        Args:
            sourc: Path to the file

        Returns:
            Extracted text
        """
        file_path = Path(sourc)
        file_ext = file_path.suffix.lower()

        if file_ext == '.pdf':
            return self._extract_text_from_pdf(str(file_path))
        elif file_ext in ['.docx', '.doc']:
            return self._extract_text_from_docx(str(file_path))
        elif file_ext == '.txt':
            # For plain text files, read and return content as-is
            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    return f.read()
            except UnicodeDecodeError:
                # Try with latin-1 encoding as fallback
                with open(file_path, 'r', encoding='latin-1') as f:
                    return f.read()
        else:
            raise ValueError(f"Unsupported file format: {file_ext}. Only PDF, DOCX, and TXT are supported.")

    def _extract_text_from_pdf(self, pdf_path: str) -> str:
        """
        Extract text from PDF using PyMuPDF and OCR.

        Args:
            pdf_path: Path to the PDF file

        Returns:
            Extracted text from all pages
        """
        try:
            import fitz  # PyMuPDF
        except ImportError as e:
            raise RuntimeError("PyMuPDF (fitz) is required. Install with: pip install pymupdf") from e
        try:
            import pytesseract
        except ImportError as e:
            raise RuntimeError("pytesseract is required for OCR. Install with: pip install pytesseract") from e
        try:
            from PIL import Image
        except ImportError as e:
            raise RuntimeError("Pillow is required for OCR. Install with: pip install Pillow") from e

        text = ""
        pdf_path = Path(pdf_path)

        with fitz.open(pdf_path) as doc:
            num_pages = len(doc)
            for page_index in range(num_pages):
                page = doc[page_index]
                # Render page to image
                pix = page.get_pixmap()
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)

                page_text = pytesseract.image_to_string(img)
                text += f"\n--- Page {page_index + 1} ---\n{page_text}"
                print(f"Processed page {page_index + 1}/{num_pages}")

        return text.strip()

    def _extract_text_from_docx(self, docx_path: str) -> str:
        """
        Extract text from DOCX file.

        Args:
            docx_path: Path to the DOCX file

        Returns:
            Extracted text from document
        """
        try:
            from docx import Document
        except ImportError as e:
            raise RuntimeError("python-docx is required. Install with: pip install python-docx") from e

        try:
            doc = Document(docx_path)

            # Extract text from paragraphs
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)

            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = ' | '.join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        full_text.append(row_text)

            return '\n'.join(full_text)

        except Exception as e:
            raise Exception(f"Error extracting text from DOCX: {str(e)}")

    def parse_contract(self, contract_text: str) -> Dict:
        """
        Parse contract text using Claude AI and return structured JSON.

        Args:
            contract_text: The extracted contract text

        Returns:
            Dictionary containing parsed contract information
        """
        prompt = f"""Analyze this contract and extract:
0. Contract type
1. Party names and roles (client vs contractor) - Payer and Receiver, change as per contract, where if I hire someone I make the contract but I pay, but if I am a landlord/provide service, I make the contract but they pay.
2. Payment SUI/crypto addresses
3. Payment amounts (Normal currency or Crypto called SUI) and schedule
4. Deadlines and milestones
5. Penalty clauses
6. Deliverables

Contract:
{contract_text}

Return as JSON with the following structure:

**Instructions:**
1. Scan the text for metadata matching the patterns above.
2. For each distinct pattern found, extract the metadata components.
3. If any specific field cannot be found within a pattern, its value should be "Unknown".

**Return valid JSON only (no markdown, no code blocks):**
{{
  "contract_type": "Service/Rent/Employment/etc",
  "parties": {{
    "client": {{
      "name": "string",
      "email": "string or null",
      "address": "string (wallet/crypto address if available) or null",
      "role": "string"
    }},
    "contractor": {{
      "name": "string",
      "email": "string or null",
      "address": "string (wallet/crypto address if available) or null",
      "role": "string"
    }},
    "payer": {{
      "name": "string",
      "address": "string (wallet/crypto address if available)",
      "role": "string"
    }},
    "receiver": {{
      "name": "string",
      "address": "string (wallet/crypto address if available)",
      "role": "string"
    }}
  }},
  "payments": {{
    "upfront": number or null,
    "completion": number or null,
    "total": number,
    "currency": "SUI/USD/etc",
    "schedule": "one-time/monthly/weekly/etc or null",
    "recurring": true/false
  }},
  "deadlines": {{
    "start_date": "YYYY-MM-DD or null",
    "final": "YYYY-MM-DD or null",
    "milestones": []
  }},
  "penalties": {{
    "late_fee_per_day": number or null,
    "max_penalty": number or null,
    "other_penalties": []
  }},
  "deliverables": []
}}

IMPORTANT: Return actual numbers (not strings) for all numeric fields like upfront, completion, total, late_fee_per_day, max_penalty."""

        try:
            # Make the API call using Claude Sonnet 4.5
            message = self.client.messages.create(
                model=self.model,
                max_tokens=4096,
                temperature=0,
                messages=[
                    {
                        "role": "user",
                        "content": prompt
                    }
                ]
            )

            # Validate response structure
            if not message.content or len(message.content) == 0:
                raise Exception("Empty response from Claude API")

            # Extract the response text
            response_text = message.content[0].text.strip()

            # Clean up the response (remove markdown code blocks if present)
            response_text = self._clean_json_response(response_text)

            # Parse JSON
            parsed_data = json.loads(response_text)

            # Convert numeric string values to actual numbers
            parsed_data = self._normalize_numeric_fields(parsed_data)

            # Add metadata
            parsed_data['_metadata'] = {
                'parsed_at': datetime.utcnow().isoformat(),
                'model': self.model
            }

            return parsed_data

        except json.JSONDecodeError as e:
            error_msg = f"Failed to parse JSON from Claude response: {str(e)}"
            if 'response_text' in locals():
                error_msg += f"\nResponse: {response_text[:500]}"
            raise Exception(error_msg)
        except Exception as e:
            raise Exception(f"Error parsing contract with Claude: {str(e)}")

    def _normalize_numeric_fields(self, data: Dict) -> Dict:
        """
        Convert numeric string values to actual numbers.

        Args:
            data: Parsed contract data

        Returns:
            Data with normalized numeric fields
        """
        # Normalize payment fields
        if 'payments' in data and isinstance(data['payments'], dict):
            for field in ['upfront', 'completion', 'total', 'amount']:
                if field in data['payments'] and data['payments'][field] is not None:
                    try:
                        if isinstance(data['payments'][field], str):
                            # Remove currency symbols and commas
                            clean_value = data['payments'][field].replace('$', '').replace(',', '').strip()
                            data['payments'][field] = float(clean_value)
                        elif isinstance(data['payments'][field], (int, float)):
                            data['payments'][field] = float(data['payments'][field])
                    except (ValueError, AttributeError):
                        pass

        # Normalize penalty fields
        if 'penalties' in data and isinstance(data['penalties'], dict):
            for field in ['late_fee_per_day', 'max_penalty']:
                if field in data['penalties'] and data['penalties'][field] is not None:
                    try:
                        if isinstance(data['penalties'][field], str):
                            # Remove currency symbols and commas
                            clean_value = data['penalties'][field].replace('$', '').replace(',', '').strip()
                            data['penalties'][field] = float(clean_value)
                        elif isinstance(data['penalties'][field], (int, float)):
                            data['penalties'][field] = float(data['penalties'][field])
                    except (ValueError, AttributeError):
                        pass

        return data

    def _clean_json_response(self, response: str) -> str:
        """
        Clean the JSON response by removing markdown code blocks and extra whitespace.

        Args:
            response: Raw response from Claude

        Returns:
            Cleaned JSON string
        """
        import re

        # Remove markdown code blocks
        response = re.sub(r'^```json\s*', '', response, flags=re.MULTILINE)
        response = re.sub(r'^```\s*', '', response, flags=re.MULTILINE)
        response = response.strip()

        # If response starts with any explanatory text before {, remove it
        json_start = response.find('{')
        if json_start > 0:
            response = response[json_start:]

        # If response ends with any explanatory text after }, remove it
        json_end = response.rfind('}')
        if json_end > 0:
            response = response[:json_end + 1]

        return response

    def parse_contract_from_file(self, file_path: str) -> Dict:
        """
        Complete workflow: Extract text from file and parse contract.

        Args:
            file_path: Path to PDF, DOCX, or TXT file

        Returns:
            Dictionary containing parsed contract information
        """
        # Extract text
        contract_text = self.extract_text(file_path)

        # Validate extracted text
        if not contract_text or len(contract_text.strip()) < 100:
            raise ValueError("Extracted text is too short. The file may be corrupted or empty.")

        # Parse contract
        parsed_data = self.parse_contract(contract_text)

        # Add file metadata
        file_path_obj = Path(file_path)
        parsed_data['_metadata']['file_path'] = file_path_obj.name
        parsed_data['_metadata']['file_size_bytes'] = file_path_obj.stat().st_size
        parsed_data['_metadata']['file_type'] = file_path_obj.suffix

        return parsed_data

    def generate_human_summary(self, parsed_data: Dict) -> str:
        """
        Generate a human-readable summary from parsed contract data.

        Args:
            parsed_data: Parsed contract dictionary

        Returns:
            Human-readable summary string
        """
        summary_lines = []
        summary_lines.append("CONTRACT SUMMARY")
        summary_lines.append("=" * 60)

        # Contract Type
        if 'contract_type' in parsed_data:
            summary_lines.append(f"\nContract Type: {parsed_data['contract_type']}")

        # Parties
        if 'parties' in parsed_data and parsed_data['parties']:
            summary_lines.append("\nPARTIES:")
            parties = parsed_data['parties']
            if 'client' in parties and parties['client']:
                client = parties['client']
                summary_lines.append(f"  Client: {client.get('name', 'N/A')} ({client.get('role', 'N/A')})")
                if client.get('email'):
                    summary_lines.append(f"    Email: {client['email']}")
            if 'contractor' in parties and parties['contractor']:
                contractor = parties['contractor']
                summary_lines.append(f"  Contractor: {contractor.get('name', 'N/A')} ({contractor.get('role', 'N/A')})")
                if contractor.get('email'):
                    summary_lines.append(f"    Email: {contractor['email']}")
            if 'payer' in parties and parties['payer']:
                payer = parties['payer']
                summary_lines.append(f"  Payer: {payer.get('name', 'N/A')}")
                if payer.get('address'):
                    summary_lines.append(f"    Address: {payer['address']}")
            if 'receiver' in parties and parties['receiver']:
                receiver = parties['receiver']
                summary_lines.append(f"  Receiver: {receiver.get('name', 'N/A')}")
                if receiver.get('address'):
                    summary_lines.append(f"    Address: {receiver['address']}")

        # Payments
        if 'payments' in parsed_data and parsed_data['payments']:
            summary_lines.append("\nPAYMENT TERMS:")
            payments = parsed_data['payments']
            currency = payments.get('currency', 'USD')
            if payments.get('upfront'):
                summary_lines.append(f"  Upfront: ${payments['upfront']:,.2f} {currency}")
            if payments.get('completion'):
                summary_lines.append(f"  On Completion: ${payments['completion']:,.2f} {currency}")
            if payments.get('total'):
                summary_lines.append(f"  Total: ${payments['total']:,.2f} {currency}")
            if payments.get('schedule'):
                summary_lines.append(f"  Schedule: {payments['schedule']}")

        # Deadlines
        if 'deadlines' in parsed_data and parsed_data['deadlines']:
            summary_lines.append("\nTIMELINE:")
            deadlines = parsed_data['deadlines']
            if deadlines.get('start_date'):
                summary_lines.append(f"  Start Date: {deadlines['start_date']}")
            if deadlines.get('final'):
                summary_lines.append(f"  Deadline: {deadlines['final']}")
            if deadlines.get('milestones') and len(deadlines['milestones']) > 0:
                summary_lines.append("  Milestones:")
                for milestone in deadlines['milestones']:
                    summary_lines.append(f"    - {milestone}")

        # Penalties
        if 'penalties' in parsed_data and parsed_data['penalties']:
            penalties = parsed_data['penalties']
            if penalties.get('late_fee_per_day'):
                summary_lines.append("\nPENALTIES:")
                summary_lines.append(f"  Late Fee: ${penalties['late_fee_per_day']:,.2f}/day")
                if penalties.get('max_penalty'):
                    summary_lines.append(f"  Maximum Penalty: ${penalties['max_penalty']:,.2f}")

        # Deliverables
        if 'deliverables' in parsed_data and parsed_data['deliverables']:
            summary_lines.append("\nDELIVERABLES:")
            for deliverable in parsed_data['deliverables']:
                summary_lines.append(f"  - {deliverable}")

        summary_lines.append("\n" + "=" * 60)

        return "\n".join(summary_lines)


# Convenience functions for direct usage
def parse_contract_file(file_path: str, api_key: Optional[str] = None) -> Dict:
    """
    Parse any supported contract file (PDF, DOCX, or TXT).

    Args:
        file_path: Path to contract file
        api_key: Optional Anthropic API key

    Returns:
        Parsed contract data as dictionary
    """
    parser = ContractParser(api_key=api_key)
    return parser.parse_contract_from_file(file_path)


def extract_text_from_file(file_path: str) -> str:
    """
    Extract text from any supported file format.

    Args:
        file_path: Path to the file

    Returns:
        Extracted text
    """
    parser = ContractParser()
    return parser.extract_text(file_path)


if __name__ == "__main__":
    # Example usage
    import sys

    if len(sys.argv) < 2:
        print("Usage: python contract_parse.py <path_to_contract.pdf|docx|txt>")
        sys.exit(1)

    file_path = sys.argv[1]

    print(f"Parsing contract: {file_path}\n")

    try:
        # Parse the contract
        result = parse_contract_file(file_path)

        print("Parsed contract data:")
        print(json.dumps(result, indent=2))

    except Exception as e:
        print(f"Error: {str(e)}")
        sys.exit(1)
